"""Import C:\\Users\\Admin\\Downloads\\KODY.xlsx into KodyTable.kt."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

XLSX = Path(r"C:\Users\Admin\Downloads\KODY.xlsx")
OUT = Path(__file__).resolve().parents[1] / "app/src/main/java/com/driver/pro/utils/KodyTable.kt"


def num(v, as_int: bool = False):
    if v is None:
        return 0 if as_int else 0.0
    if isinstance(v, str) and v.strip() in ("", "?"):
        return 0 if as_int else 0.0
    try:
        f = float(v)
        return int(round(f)) if as_int else f
    except (TypeError, ValueError):
        return 0 if as_int else 0.0


def fmt_miles(a: float) -> str:
    if a == int(a):
        return f"{int(a)}.0"
    return f"{a:.4f}".rstrip("0").rstrip(".")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    patterns: list[tuple[int, ...]] = []
    pat_index: dict[tuple[int, ...], int] = {}
    rows: list[tuple[str, int, float, int]] = []
    for r in range(2, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if not raw:
            continue
        code = re.sub(r"[^A-Z0-9]", "", str(raw).strip().upper())
        if not code:
            continue
        hours = tuple(int(num(ws.cell(r, c).value, True)) for c in range(2, 26))
        if hours not in pat_index:
            pat_index[hours] = len(patterns)
            patterns.append(hours)
        a = num(ws.cell(r, 26).value, False)
        b = num(ws.cell(r, 27).value, True)
        rows.append((code, pat_index[hours], a, b))

    lines = [
        "package com.driver.pro.utils",
        "",
        "/** Compact import of KODY.xlsx: X1–X24 hour points, A miles, B minutes. */",
        "internal data class KodyRowRef(val pattern: Int, val returnMiles: Double, val returnMinutes: Int)",
        "",
        "internal object KodyTable {",
        "    val patterns: Array<IntArray> = arrayOf(",
    ]
    for p in patterns:
        lines.append("        intArrayOf(" + ", ".join(str(x) for x in p) + "),")
    chunk = 400
    n_chunks = (len(rows) + chunk - 1) // chunk
    lines += [
        "    )",
        "",
        "    val rows: Map<String, KodyRowRef> = HashMap<String, KodyRowRef>(2832).apply {",
    ]
    for i in range(n_chunks):
        lines.append(f"        load{i}(this)")
    lines.append("    }")
    lines.append("")
    for i in range(n_chunks):
        part = rows[i * chunk : (i + 1) * chunk]
        lines.append(
            "    private fun load"
            f"{i}(into: MutableMap<String, KodyRowRef>) {{"
        )
        for code, pi, a, b in part:
            lines.append(
                f'        into.put("{code}", KodyRowRef({pi}, {fmt_miles(a)}, {b}))'
            )
        lines.append("    }")
        lines.append("")
    lines += [
        "}",
        "",
    ]
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {OUT} patterns={len(patterns)} rows={len(rows)} bytes={OUT.stat().st_size}")


if __name__ == "__main__":
    main()
